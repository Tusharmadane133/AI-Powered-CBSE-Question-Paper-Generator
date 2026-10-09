"""Word (.docx) export of a generated paper (question paper + answer key)."""
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

INSTRUCTIONS = ["All questions are compulsory.",
                "Marks are indicated against each section.",
                "Write neat and legible answers."]


def _setup(doc):
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(11)
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    rf.set(qn("w:ascii"), "Calibri")
    rf.set(qn("w:hAnsi"), "Calibri")
    rf.set(qn("w:cs"), "Mangal")        # Devanagari (Hindi / Marathi / Sanskrit)
    rf.set(qn("w:eastAsia"), "Mangal")
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2)
        s.top_margin = s.bottom_margin = Cm(1.8)


def _center(doc, text, size=12, bold=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold, r.font.size = bold, Pt(size)
    p.paragraph_format.space_after = Pt(2)
    return p


def build_docx(paper) -> bytes:
    m = paper["meta"]
    doc = Document()
    _setup(doc)
    if m.get("school"):
        _center(doc, m["school"], 16, True)
    _center(doc, m["title"], 14, True)
    _center(doc, f"Class: {m['class']}      Subject: {m['subject']}", 12)
    _center(doc, f"Time: {m['duration']} minutes      Max. Marks: {m['total_marks']}", 12)
    doc.add_paragraph("_" * 90).alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    p.add_run("General Instructions:").bold = True
    for i, t in enumerate(INSTRUCTIONS, 1):
        doc.add_paragraph(f"{i}. {t}").paragraph_format.space_after = Pt(0)

    for sec in paper["sections"]:
        total = sec["marks_each"] * len(sec["questions"])
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        r = h.add_run(f"{sec['name']} — {sec['label']}  ({sec['marks_each']} mark(s) each, {total} marks)")
        r.bold, r.font.size = True, Pt(12)
        for q in sec["questions"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            p.add_run(f"Q{q['no']}. ").bold = True
            p.add_run(q["question"])
            p.add_run(f"  [{sec['marks_each']}]").italic = True
            for L, o in zip("ABCD", q.get("options", [])):
                op = doc.add_paragraph(f"({L}) {o}")
                op.paragraph_format.left_indent = Cm(1)
                op.paragraph_format.space_after = Pt(0)

    doc.add_page_break()
    _center(doc, "ANSWER KEY / MARKING SCHEME", 14, True)
    for sec in paper["sections"]:
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(10)
        h.add_run(f"{sec['name']} — {sec['label']}").bold = True
        for q in sec["questions"]:
            p = doc.add_paragraph()
            p.add_run(f"Q{q['no']}. ").bold = True
            p.add_run(q["answer"])
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
